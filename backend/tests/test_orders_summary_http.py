"""HTTP contract for the day summary, place, and export (no network)."""

import csv
import io
from datetime import date, timedelta
from unittest.mock import patch

import bcrypt
from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import func, select

from app.catalog.models import Dish
from app.db import SessionLocal
from app.orders.export_render import CSV_COLUMNS, EMPLOYEE_SHEET_NAME
from app.orders.models import OfficeOrder, OfficeOrderStatus
from app.orders.service import DAY_STILL_COLLECTING_DETAIL
from app.providers import MealtyError, NormalizedDish
from app.users.models import Department, User, UserRole
from tests.conftest import BearerAuth

DAY = date.today() + timedelta(days=18)


class _FailingProvider:
    source = "mealty"

    async def fetch_catalog(self) -> list[NormalizedDish]:
        raise MealtyError("Mealty request failed")


def _order_path(day: date = DAY) -> str:
    return f"/api/orders/{day.isoformat()}"


async def _add_dish(external_id: str, **overrides: object) -> int:
    values: dict[str, object] = {
        "source": "mealty",
        "external_id": external_id,
        "seller_product_id": None,
        "name": f"Блюдо {external_id}",
        "subtitle": "",
        "description": None,
        "category": "main_dish",
        "price_kopecks": 41000,
        "old_price_kopecks": None,
        "weight_g": 300,
        "proteins": None,
        "fats": None,
        "carbs": None,
        "calories": None,
        "image_url": None,
        "available": True,
    }
    values.update(overrides)
    async with SessionLocal() as session:
        dish = Dish(**values)
        session.add(dish)
        await session.commit()
        await session.refresh(dish)
        return dish.id


async def _employee(
    client: AsyncClient,
    email: str,
    full_name: str,
    department_name: str | None = None,
) -> BearerAuth:
    password_hash = bcrypt.hashpw(b"secret", bcrypt.gensalt()).decode("utf-8")
    async with SessionLocal() as session:
        department_id = None
        if department_name is not None:
            department = Department(name=department_name)
            session.add(department)
            await session.flush()
            department_id = department.id
        user = User(
            email=email,
            password_hash=password_hash,
            full_name=full_name,
            role=UserRole.EMPLOYEE,
            department_id=department_id,
            is_active=True,
        )
        session.add(user)
        await session.commit()
        await session.refresh(user)
        user_id = user.id
    response = await client.post(
        "/api/auth/login",
        json={"email": email, "password": "secret"},
    )
    assert response.status_code == 200, response.text
    return BearerAuth(
        token=response.json()["access_token"],
        user_id=user_id,
        role=UserRole.EMPLOYEE,
        email=email,
    )


async def _procurement_user(client: AsyncClient, email: str) -> BearerAuth:
    password_hash = bcrypt.hashpw(b"secret", bcrypt.gensalt()).decode("utf-8")
    async with SessionLocal() as session:
        user = User(
            email=email,
            password_hash=password_hash,
            full_name="Вторая Закупка",
            role=UserRole.PROCUREMENT,
            is_active=True,
        )
        session.add(user)
        await session.commit()
        await session.refresh(user)
        user_id = user.id
    response = await client.post(
        "/api/auth/login",
        json={"email": email, "password": "secret"},
    )
    assert response.status_code == 200, response.text
    return BearerAuth(
        token=response.json()["access_token"],
        user_id=user_id,
        role=UserRole.PROCUREMENT,
        email=email,
    )


async def _save(
    client: AsyncClient,
    auth: BearerAuth,
    dish_id: int,
    qty: int = 1,
    day: date = DAY,
) -> None:
    saved = await client.put(
        f"/api/plans/{day.isoformat()}",
        headers=auth.headers,
        json={"items": [{"dish_id": dish_id, "qty": qty}]},
    )
    assert saved.status_code == 200, saved.text


async def _cutoff(
    client: AsyncClient,
    auth: BearerAuth,
    day: date = DAY,
) -> None:
    with patch("app.catalog.service.MealtyProvider", _FailingProvider):
        closed = await client.post(
            f"/api/orders/{day.isoformat()}/cutoff",
            headers=auth.headers,
        )
    assert closed.status_code == 200, closed.text


async def _order_status(day: date = DAY) -> OfficeOrderStatus | None:
    async with SessionLocal() as session:
        order = await session.scalar(
            select(OfficeOrder).where(OfficeOrder.delivery_date == day)
        )
        return None if order is None else order.status


async def _order_count() -> int:
    async with SessionLocal() as session:
        count = await session.scalar(select(func.count()).select_from(OfficeOrder))
        return int(count or 0)


async def test_summary_access(
    client: AsyncClient,
    employee_auth: BearerAuth,
    procurement_auth: BearerAuth,
    admin_auth: BearerAuth,
) -> None:
    missing = await client.get(_order_path())
    assert missing.status_code == 401

    denied = await client.get(_order_path(), headers=employee_auth.headers)
    assert denied.status_code == 403
    assert denied.json()["detail"] == "Forbidden"

    for auth in (procurement_auth, admin_auth):
        allowed = await client.get(_order_path(), headers=auth.headers)
        assert allowed.status_code == 200, allowed.text
        body = allowed.json()
        assert body["status"] == "collecting"
        assert body["totals"]["actual_kopecks"] is None
        assert body["totals"]["delta_kopecks"] is None
        assert body["by_employee"] == []
    assert await _order_count() == 0


async def test_summary_collecting_lists_draft_plans(
    client: AsyncClient,
    procurement_auth: BearerAuth,
) -> None:
    dish_id = await _add_dish("http-open", price_kopecks=41000)
    ivan = await _employee(
        client,
        "ivan-http@summary.test",
        "Иван Сотрудник",
        "Разработка",
    )
    maria = await _employee(client, "maria-http@summary.test", "Мария Сотрудница", "HR")
    await _save(client, ivan, dish_id, qty=2)
    await _save(client, maria, dish_id, qty=1)
    async with SessionLocal() as session:
        dish = await session.get(Dish, dish_id)
        assert dish is not None
        dish.price_kopecks = 42500
        await session.commit()

    response = await client.get(_order_path(), headers=procurement_auth.headers)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["status"] == "collecting"
    assert body["totals"]["planned_kopecks"] == 3 * 41000
    assert body["totals"]["actual_kopecks"] is None
    assert body["items"][0]["qty"] == 3
    assert body["items"][0]["unit_price_kopecks"] == 42500
    assert body["items"][0]["line_total_kopecks"] == 3 * 42500
    assert {row["full_name"] for row in body["by_employee"]} == {
        "Иван Сотрудник",
        "Мария Сотрудница",
    }
    assert all(row["actual_total_kopecks"] is None for row in body["by_employee"])
    assert await _order_count() == 0


async def test_place_access_and_idempotent_first_write(
    client: AsyncClient,
    employee_auth: BearerAuth,
    procurement_auth: BearerAuth,
    admin_auth: BearerAuth,
) -> None:
    missing = await client.post(f"{_order_path()}/place")
    assert missing.status_code == 401

    denied = await client.post(
        f"{_order_path()}/place",
        headers=employee_auth.headers,
    )
    assert denied.status_code == 403
    assert denied.json()["detail"] == "Forbidden"

    collecting = await client.post(
        f"{_order_path()}/place",
        headers=procurement_auth.headers,
    )
    assert collecting.status_code == 403
    assert collecting.json()["detail"] == DAY_STILL_COLLECTING_DETAIL
    assert await _order_count() == 0

    dish_id = await _add_dish("http-place")
    await _save(client, employee_auth, dish_id)
    await _cutoff(client, procurement_auth)

    first = await client.post(
        f"{_order_path()}/place",
        headers=procurement_auth.headers,
    )
    assert first.status_code == 200, first.text
    first_body = first.json()
    assert first_body["status"] == "placed"
    assert first_body["already_placed"] is False
    assert first_body["placed_by_user_id"] == procurement_auth.user_id
    assert first_body["delivery_date"] == DAY.isoformat()

    other = await _procurement_user(client, "other-proc@summary.test")
    second = await client.post(f"{_order_path()}/place", headers=other.headers)
    assert second.status_code == 200, second.text
    second_body = second.json()
    assert second_body["already_placed"] is True
    assert second_body["placed_at"] == first_body["placed_at"]
    assert second_body["placed_by_user_id"] == procurement_auth.user_id
    assert second_body["status"] == "placed"

    admin_repeat = await client.post(
        f"{_order_path()}/place",
        headers=admin_auth.headers,
    )
    assert admin_repeat.status_code == 200
    assert admin_repeat.json()["placed_by_user_id"] == procurement_auth.user_id

    async with SessionLocal() as session:
        order = await session.scalar(
            select(OfficeOrder).where(OfficeOrder.delivery_date == DAY)
        )
        assert order is not None
        assert order.status == OfficeOrderStatus.PLACED
        assert order.placed_by_user_id == procurement_auth.user_id
        assert order.placed_at is not None
        assert order.placed_at.tzinfo is not None


async def test_export_headers_status_and_columns(
    client: AsyncClient,
    employee_auth: BearerAuth,
    procurement_auth: BearerAuth,
) -> None:
    collecting = await client.get(
        f"{_order_path()}/export.csv",
        headers=procurement_auth.headers,
    )
    assert collecting.status_code == 403
    assert collecting.json()["detail"] == DAY_STILL_COLLECTING_DETAIL
    assert await _order_count() == 0

    denied = await client.get(
        f"{_order_path()}/export.xlsx",
        headers=employee_auth.headers,
    )
    assert denied.status_code == 403

    dish_id = await _add_dish("http-export", name="Ньокки с грибами")
    ivan = await _employee(client, "export-ivan@summary.test", "Иван Сотрудник")
    maria = await _employee(client, "export-maria@summary.test", "Мария Сотрудница")
    await _save(client, ivan, dish_id, qty=2)
    await _save(client, maria, dish_id, qty=2)
    await _cutoff(client, procurement_auth)

    csv_response = await client.get(
        f"{_order_path()}/export.csv",
        headers=procurement_auth.headers,
    )
    assert csv_response.status_code == 200, csv_response.text
    content_type = csv_response.headers["content-type"]
    assert "text/csv" in content_type
    assert "charset=utf-8" in content_type.lower()
    disposition = csv_response.headers["content-disposition"]
    assert disposition.startswith("attachment;")
    assert f"corplunch-{DAY.isoformat()}.csv" in disposition
    assert csv_response.content.startswith(b"\xef\xbb\xbf")
    rows = list(csv.reader(io.StringIO(csv_response.content.decode("utf-8-sig"))))
    assert rows[0] == list(CSV_COLUMNS)
    assert rows[1][0] == "http-export"
    assert rows[1][1] == "Ньокки с грибами"
    assert rows[1][2] == "4"
    assert await _order_status() == OfficeOrderStatus.EXPORTED

    again = await client.get(
        f"{_order_path()}/export.csv",
        headers=procurement_auth.headers,
    )
    assert again.status_code == 200
    assert await _order_status() == OfficeOrderStatus.EXPORTED

    xlsx_response = await client.get(
        f"{_order_path()}/export.xlsx",
        headers=procurement_auth.headers,
    )
    assert xlsx_response.status_code == 200, xlsx_response.text
    assert (
        xlsx_response.headers["content-type"]
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert (
        f"corplunch-{DAY.isoformat()}.xlsx"
        in xlsx_response.headers["content-disposition"]
    )
    workbook = load_workbook(io.BytesIO(xlsx_response.content))
    assert EMPLOYEE_SHEET_NAME in workbook.sheetnames
    people = workbook[EMPLOYEE_SHEET_NAME]
    names = {people["B2"].value, people["B3"].value}
    assert names == {"Иван Сотрудник", "Мария Сотрудница"}
    assert await _order_status() == OfficeOrderStatus.EXPORTED


async def test_export_after_place_keeps_placed(
    client: AsyncClient,
    employee_auth: BearerAuth,
    procurement_auth: BearerAuth,
) -> None:
    dish_id = await _add_dish("http-placed-export")
    await _save(client, employee_auth, dish_id)
    await _cutoff(client, procurement_auth)
    placed = await client.post(
        f"{_order_path()}/place",
        headers=procurement_auth.headers,
    )
    assert placed.status_code == 200, placed.text

    downloaded = await client.get(
        f"{_order_path()}/export.csv",
        headers=procurement_auth.headers,
    )
    assert downloaded.status_code == 200, downloaded.text
    assert downloaded.content.startswith(b"\xef\xbb\xbf")
    assert await _order_status() == OfficeOrderStatus.PLACED

    summary = await client.get(_order_path(), headers=procurement_auth.headers)
    assert summary.status_code == 200
    assert summary.json()["status"] == "placed"


async def test_procurement_scenario_two_employees_file_and_place(
    client: AsyncClient,
    procurement_auth: BearerAuth,
) -> None:
    """roles.md procurement steps 5–7: summary, file, safe repeat place."""
    gnocchi = await _add_dish("875", name="Картофельные ньокки", price_kopecks=41000)
    salad = await _add_dish("902", name="Салат", price_kopecks=18000)
    ivan = await _employee(
        client,
        "scenario-ivan@summary.test",
        "Иван Сотрудник",
        "Разработка",
    )
    maria = await _employee(
        client,
        "scenario-maria@summary.test",
        "Мария Сотрудница",
        "HR",
    )
    ivan_plan = await client.put(
        f"/api/plans/{DAY.isoformat()}",
        headers=ivan.headers,
        json={"items": [{"dish_id": gnocchi, "qty": 2}]},
    )
    maria_plan = await client.put(
        f"/api/plans/{DAY.isoformat()}",
        headers=maria.headers,
        json={
            "items": [
                {"dish_id": gnocchi, "qty": 1},
                {"dish_id": salad, "qty": 1},
            ]
        },
    )
    assert ivan_plan.status_code == 200, ivan_plan.text
    assert maria_plan.status_code == 200, maria_plan.text
    await _cutoff(client, procurement_auth)

    summary = await client.get(_order_path(), headers=procurement_auth.headers)
    assert summary.status_code == 200, summary.text
    body = summary.json()
    assert {row["full_name"] for row in body["by_employee"]} == {
        "Иван Сотрудник",
        "Мария Сотрудница",
    }
    qty = {item["external_id"]: item["qty"] for item in body["items"]}
    assert qty == {"875": 3, "902": 1}
    assert all(isinstance(item["line_total_kopecks"], int) for item in body["items"])

    csv_file = await client.get(
        f"{_order_path()}/export.csv",
        headers=procurement_auth.headers,
    )
    xlsx_file = await client.get(
        f"{_order_path()}/export.xlsx",
        headers=procurement_auth.headers,
    )
    assert csv_file.status_code == 200
    assert xlsx_file.status_code == 200
    parsed = list(csv.DictReader(io.StringIO(csv_file.content.decode("utf-8-sig"))))
    assert {row["external_id"] for row in parsed} == {"875", "902"}
    book = load_workbook(io.BytesIO(xlsx_file.content))
    assert book.sheetnames[0] == "Блюда"
    assert EMPLOYEE_SHEET_NAME in book.sheetnames

    first = await client.post(
        f"{_order_path()}/place",
        headers=procurement_auth.headers,
    )
    second = await client.post(
        f"{_order_path()}/place",
        headers=procurement_auth.headers,
    )
    assert first.status_code == 200, first.text
    assert second.status_code == 200, second.text
    assert second.json()["already_placed"] is True
    assert second.json()["placed_at"] == first.json()["placed_at"]
    assert second.json()["placed_by_user_id"] == procurement_auth.user_id
    assert await _order_status() == OfficeOrderStatus.PLACED
