"""CSV/XLSX render from fixture DTO (no live orders aggregate)."""

import csv
import io
import json
from pathlib import Path

from openpyxl import load_workbook

from app.orders.export_render import (
    CSV_COLUMNS,
    EMPLOYEE_COLUMNS,
    EMPLOYEE_SHEET_NAME,
    render_summary_csv,
    render_summary_xlsx,
)

FIXTURE = Path(__file__).parent / "fixtures" / "order_summary.json"


def _summary() -> dict:
    return json.loads(FIXTURE.read_text(encoding="utf-8"))


def test_render_summary_csv_columns_and_rows() -> None:
    raw = render_summary_csv(_summary())
    text = raw.decode("utf-8-sig")
    rows = list(csv.DictReader(io.StringIO(text)))
    assert list(rows[0].keys()) == list(CSV_COLUMNS)
    assert len(rows) == 2
    assert rows[0]["external_id"] == "875"
    assert rows[0]["name"] == "Картофельные ньокки с грибами"
    assert rows[0]["qty"] == "4"
    assert rows[0]["unit_price_kopecks"] == "42500"
    assert rows[0]["line_total_kopecks"] == "170000"
    assert rows[1]["external_id"] == "902"


def test_render_summary_xlsx_has_dishes_and_employees() -> None:
    raw = render_summary_xlsx(_summary())
    workbook = load_workbook(io.BytesIO(raw))
    assert workbook.sheetnames[0] == "Блюда"
    assert EMPLOYEE_SHEET_NAME in workbook.sheetnames

    dishes = workbook["Блюда"]
    assert [cell.value for cell in dishes[1]] == list(CSV_COLUMNS)
    assert dishes["A2"].value == "875"
    assert dishes["C2"].value == 4
    assert dishes["E2"].value == 170000
    assert dishes["A3"].value == "902"

    people = workbook[EMPLOYEE_SHEET_NAME]
    assert [cell.value for cell in people[1]] == list(EMPLOYEE_COLUMNS)
    assert people["B2"].value == "Иван Сотрудник"
    assert people["E2"].value == 85000
    assert people["B3"].value == "Мария Сотрудница"


def test_empty_items_still_emits_header() -> None:
    raw = render_summary_csv({"items": [], "by_employee": []})
    text = raw.decode("utf-8-sig").strip()
    assert text == ",".join(CSV_COLUMNS)
